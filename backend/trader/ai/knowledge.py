"""Turning uploaded files into searchable text: extraction, chunking and a small BM25 ranker.

Uploaded content is untrusted data. It is only ever shown to the model as quoted material, never as instructions.
"""

from __future__ import annotations

import csv
import io
import math
import re
import shutil
from collections import Counter
from typing import Callable

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 400_000
MAX_PDF_PAGES = 600
MIN_PAGE_CHARS = 25  # a page with less text than this is treated as a scan
MAX_OCR_PAGES = 150
OCR_MAX_SIDE = 2000  # downscale big scans before OCR: faster, and plenty for printed text
CHUNK_CHARS = 1200
CHUNK_OVERLAP = 150

IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}
TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml", ".py", ".log", ".xml", ".html", ".htm", ".ini", ".toml"}


class UnsupportedFile(ValueError):
    pass


def _ext(filename: str) -> str:
    return "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def kind_of(filename: str, content_type: str = "") -> str:
    ext = _ext(filename)
    if ext == ".pdf" or content_type == "application/pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext in (".xlsx", ".xlsm"):
        return "xlsx"
    if ext in IMAGE_TYPES or content_type.startswith("image/"):
        return "image"
    if ext in TEXT_EXT or content_type.startswith("text/") or content_type in ("application/json", "application/xml"):
        return "text"
    raise UnsupportedFile(f"Can't read {filename!r}. Supported: PDF, Word (.docx), Excel (.xlsx), images, and text files "
                          "(txt, md, csv, json, ...). Old .doc/.xls files: save them as .docx/.xlsx first.")


def image_mime(filename: str, content_type: str = "") -> str:
    return IMAGE_TYPES.get(_ext(filename)) or (content_type if content_type.startswith("image/") else "image/png")


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise UnsupportedFile("This PDF is password protected.")
    parts: list[str] = []
    total = 0
    for i, page in enumerate(reader.pages, 1):  # one page at a time; stop once enough text is kept (small server, big books)
        if i > MAX_PDF_PAGES or total >= MAX_TEXT_CHARS:
            break
        text = (page.extract_text() or "").strip()
        total += len(text)
        parts.append(f"[page {i}]\n{text}")
    return "\n\n".join(parts)


_PAGE_MARK = re.compile(r"^\[page (\d+)\]\n", re.M)


def blank_pdf_pages(text: str) -> list[int]:
    """Page numbers whose extracted text is (nearly) empty: scanned pages that OCR could read."""
    marks = list(_PAGE_MARK.finditer(text))
    out = []
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        if len(text[m.end():end].strip()) < MIN_PAGE_CHARS:
            out.append(int(m.group(1)))
    return out


def merge_page_text(text: str, found: dict[int, str]) -> str:
    """Put OCR text into the matching '[page N]' sections of an extracted PDF text."""
    marks = list(_PAGE_MARK.finditer(text))
    parts = []
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        n = int(m.group(1))
        body = found.get(n) or text[m.end():end].strip()
        parts.append(f"[page {n}]\n{body}")
    return "\n\n".join(parts)


def ocr_available() -> bool:
    """OCR needs the tesseract program (apt install tesseract-ocr) and the pytesseract package."""
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return shutil.which("tesseract") is not None


def _tesseract(image) -> str:
    import pytesseract

    if max(image.size) > OCR_MAX_SIDE:
        image.thumbnail((OCR_MAX_SIDE, OCR_MAX_SIDE))
    return pytesseract.image_to_string(image.convert("L"), timeout=90)


def ocr_pdf_pages(data: bytes, pages: list[int], ocr: Callable | None = None) -> dict[int, str]:
    """OCR the pictures on the given (1-based) pages. A page that fails is skipped, never fatal."""
    from pypdf import PdfReader

    ocr = ocr or _tesseract
    reader = PdfReader(io.BytesIO(data))
    found: dict[int, str] = {}
    for n in pages:
        try:
            texts = [ocr(img.image).strip() for img in reader.pages[n - 1].images]
        except Exception:  # noqa: BLE001 - odd image, tesseract timeout...
            continue
        if any(texts):
            found[n] = "\n".join(t for t in texts if t)
    return found


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        parts += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"[sheet {ws.title}]")
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= 2000:
                out.append("... (rows truncated)")
                break
            if any(c is not None for c in row):
                out.append(" | ".join("" if c is None else str(c) for c in row))
    return "\n".join(out)


def _csv(data: bytes) -> str:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        return text
    return "\n".join(" | ".join(r) for _, r in zip(range(3000), csv.reader(io.StringIO(text), dialect)))


def extract_text(filename: str, content_type: str, data: bytes) -> tuple[str, str]:
    """Return (kind, text). Images return empty text: they are read by a vision model, see ai/providers."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise UnsupportedFile(f"{filename!r} is larger than {MAX_UPLOAD_BYTES // 1024 // 1024} MB.")
    kind = kind_of(filename, content_type)
    try:
        if kind == "pdf":
            text = _pdf(data)
        elif kind == "docx":
            text = _docx(data)
        elif kind == "xlsx":
            text = _xlsx(data)
        elif kind == "image":
            return kind, ""
        else:
            text = _csv(data) if _ext(filename) in (".csv", ".tsv") else _decode(data)
    except UnsupportedFile:
        raise
    except Exception as exc:  # noqa: BLE001 - corrupt file: report, don't 500
        raise UnsupportedFile(f"Could not read {filename!r}: {exc}") from exc
    return kind, re.sub(r"\n{3,}", "\n\n", text).strip()[:MAX_TEXT_CHARS]


def chunk_text(text: str) -> list[str]:
    """Split on paragraph boundaries into ~CHUNK_CHARS pieces with a little overlap."""
    chunks: list[str] = []
    cur = ""
    for para in re.split(r"\n\s*\n", text):
        while len(para) > CHUNK_CHARS:  # one huge paragraph (e.g. a table dump)
            if cur:
                chunks.append(cur)
                cur = ""
            chunks.append(para[:CHUNK_CHARS])
            para = para[CHUNK_CHARS - CHUNK_OVERLAP:]
        if len(cur) + len(para) + 2 > CHUNK_CHARS and cur:
            chunks.append(cur)
            cur = cur[-CHUNK_OVERLAP:]
        cur = f"{cur}\n\n{para}".strip() if cur else para
    if cur.strip():
        chunks.append(cur)
    return chunks


_WORD = re.compile(r"[a-z0-9_.%]+")


def _tokens(s: str) -> list[str]:
    return _WORD.findall(s.lower())


def bm25_rank(query: str, chunks: list[str], k: int = 5, k1: float = 1.5, b: float = 0.75) -> list[tuple[int, float]]:
    """(index, score) of the best chunks for the query; chunks sharing no query word are dropped."""
    q = set(_tokens(query))
    docs = [Counter(_tokens(c)) for c in chunks]
    if not q or not docs:
        return []
    avg = sum(sum(d.values()) for d in docs) / len(docs) or 1.0
    df = {t: sum(1 for d in docs if t in d) for t in q}
    scored = []
    for i, d in enumerate(docs):
        length = sum(d.values())
        s = 0.0
        for t in q:
            if df[t] and d.get(t):
                idf = math.log(1 + (len(docs) - df[t] + 0.5) / (df[t] + 0.5))
                s += idf * d[t] * (k1 + 1) / (d[t] + k1 * (1 - b + b * length / avg))
        if s > 0:
            scored.append((i, s))
    return sorted(scored, key=lambda x: -x[1])[:k]
