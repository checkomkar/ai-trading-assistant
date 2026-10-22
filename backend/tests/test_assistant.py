import io
import json

import pytest
from fastapi.testclient import TestClient

from trader.ai import chat as chat_mod
from trader.ai import knowledge as kn
from trader.ai.providers import LLMReply, TranslatorUnavailable
from trader.api.main import create_app
from trader.config import Settings


# ---- extraction / chunking / search --------------------------------------------------------------
def _docx(text: str) -> bytes:
    from docx import Document

    d = Document()
    d.add_paragraph(text)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "stop", "20%"
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Journal"
    ws.append(["date", "pnl"])
    ws.append(["2026-10-20", 1250])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _pdf(text: str) -> bytes:
    """A minimal one-page PDF with a text line (enough for pypdf to extract)."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 100] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"]
    stream = f"BT /F1 12 Tf 10 50 Td ({text}) Tj ET"
    objs += [f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream", "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    x = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    return out + f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF".encode()


@pytest.mark.parametrize("name,ctype,data,expect", [
    ("rules.txt", "text/plain", b"Buy calls when RSI > 60", "RSI > 60"),
    ("trades.csv", "text/csv", b"date,pnl\n2026-10-20,1250\n", "2026-10-20 | 1250"),
    ("plan.docx", "", _docx("Opening range breakout plan"), "Opening range breakout plan"),
    ("journal.xlsx", "", _xlsx(), "[sheet Journal]"),
    ("notes.pdf", "application/pdf", _pdf("Trail stop after thirty percent"), "Trail stop after thirty percent"),
], ids=["txt", "csv", "docx", "xlsx", "pdf"])
def test_extract_text(name, ctype, data, expect):
    kind, text = kn.extract_text(name, ctype, data)
    assert expect in text and kind in ("text", "docx", "xlsx", "pdf")


def test_extract_rejections_and_images():
    assert kn.extract_text("chart.png", "image/png", b"\x89PNG")[0] == "image"
    with pytest.raises(kn.UnsupportedFile, match="docx"):
        kn.extract_text("old.doc", "application/msword", b"x")
    with pytest.raises(kn.UnsupportedFile, match="Could not read"):
        kn.extract_text("bad.docx", "", b"not a zip")
    with pytest.raises(kn.UnsupportedFile, match="larger"):
        kn.extract_text("big.txt", "", b"x" * (kn.MAX_UPLOAD_BYTES + 1))


def test_chunking_and_bm25_ranking():
    text = "\n\n".join(["Opening range breakout uses the first fifteen minutes."] * 3
                       + ["Risk: never risk more than two percent per trade."]
                       + ["Gap-down days: skip the first candle entirely."] * 3)
    chunks = kn.chunk_text(text * 6)
    assert len(chunks) > 1 and all(len(c) <= kn.CHUNK_CHARS + kn.CHUNK_OVERLAP + 10 for c in chunks)
    ranked = kn.bm25_rank("risk percent per trade", ["opening range", "never risk two percent per trade", "gap down"])
    assert ranked[0][0] == 1 and len(ranked) == 1  # chunks with no query word are dropped


# ---- chat engine ---------------------------------------------------------------------------------
class Scripted:
    """A provider that replays queued actions and records what it was sent."""

    name, model = "fake", "m"
    label = "fake/m"

    def __init__(self, *actions):
        self.actions, self.calls, self.images = list(actions), [], []

    def complete(self, system, messages, schema, tools=None):
        self.calls.append((system, list(messages)))
        return LLMReply(json.dumps(self.actions.pop(0)))

    def describe_image(self, data, mime, prompt):
        self.images.append(mime)
        return "NIFTY 5m chart, uptrend above the 21 EMA"


def tool(name, args="{}"):
    return {"action": "tool", "tool": name, "args_json": args, "text": "", "remember": []}


def answer(text, remember=()):
    return {"action": "answer", "tool": "", "args_json": "{}", "text": text, "remember": list(remember)}


async def _noop(_):
    return {"ok": True}


async def test_engine_runs_tool_then_answers_and_keeps_cards():
    async def drafts(a):
        return {"name": "X", "_card": {"type": "strategy_draft", "name": "X"}}

    p = Scripted(tool("draft_strategy", '{"description":"x"}'), answer("Here is a draft.", ["Trades NIFTY only"]))
    r = await chat_mod.run_chat(p, {"draft_strategy": chat_mod.Tool("d", drafts)}, [], "make one", [], [])
    assert r.text == "Here is a draft." and r.cards == [{"type": "strategy_draft", "name": "X"}]
    assert r.remember == ["Trades NIFTY only"] and r.tools_used == ["draft_strategy"]
    assert "_card" not in p.calls[1][1][-1]["content"]  # the model sees the result, not the card payload


async def test_engine_survives_unknown_tool_and_tool_errors():
    async def boom(_):
        raise RuntimeError("db down")

    p = Scripted(tool("place_order"), tool("boom"), answer("ok"))
    r = await chat_mod.run_chat(p, {"boom": chat_mod.Tool("b", boom)}, [], "hi", [], [])
    assert r.text == "ok" and r.tools_used == []
    assert "unknown tool" in p.calls[1][1][-1]["content"] and "db down" in p.calls[2][1][-1]["content"]


def test_blank_pages_are_found_and_ocr_text_is_merged_back_in_place():
    from trader.ai import knowledge as k

    text = "[page 1]\nReal text on page one is long enough.\n\n[page 2]\n\n\n[page 3]\n\n\n[page 4]\nAnother page with real text here."
    assert k.blank_pdf_pages(text) == [2, 3]
    merged = k.merge_page_text(text, {2: "Scanned rule: buy the pullback"})
    assert "[page 2]\nScanned rule: buy the pullback" in merged and "Real text on page one" in merged and "Another page" in merged
    assert k.blank_pdf_pages(merged) == [3]


def test_ocr_reads_images_on_the_requested_pages_and_survives_a_bad_page(monkeypatch):
    import pypdf

    from trader.ai import knowledge as k

    class Img:
        def __init__(self, image):
            self.image = image

    class Page:
        def __init__(self, images):
            self.images = images

    class Reader:
        def __init__(self, _):
            self.pages = [Page([Img("a")]), Page([Img("bad")]), Page([Img("c")])]

    monkeypatch.setattr(pypdf, "PdfReader", Reader)

    def fake(img):
        if img == "bad":
            raise RuntimeError("tesseract timeout")
        return f" text-{img} "

    assert k.ocr_pdf_pages(b"x", [1, 2, 3], ocr=fake) == {1: "text-a", 3: "text-c"}


def test_scanned_pdf_is_stored_at_once_then_filled_in_by_ocr(client, monkeypatch):
    import time

    from trader.api.routes import chat as chat_routes

    monkeypatch.setattr(chat_routes, "ocr_available", lambda: True)
    monkeypatch.setattr(chat_routes, "extract_text", lambda f, c, d: ("pdf", "[page 1]\nA page with plenty of real text.\n\n[page 2]\n"))
    monkeypatch.setattr(chat_routes, "ocr_pdf_pages", lambda data, pages: {p: f"Scanned rule {p}: wait for the pullback" for p in pages})
    doc = upload(client, "book.pdf", b"%PDF-fake").json()
    assert doc["status"] == "ocr"
    for _ in range(50):
        docs = client.get("/knowledge/docs").json()["docs"]
        if docs[0]["status"] != "ocr":
            break
        time.sleep(0.1)
    assert docs[0]["status"] == "ready" and docs[0]["chars"] > doc["chars"]  # the scanned page's text was added


async def test_engine_asks_for_a_final_answer_instead_of_running_out_of_steps():
    p = Scripted(*[tool("t")] * (chat_mod.MAX_STEPS - 1), answer("Summary of the files."))
    r = await chat_mod.run_chat(p, {"t": chat_mod.Tool("t", _noop)}, [], "q", [], [])
    assert r.text == "Summary of the files."
    assert chat_mod.LAST_CALL in p.calls[-1][1][-1]["content"]


async def test_engine_retries_once_when_the_model_returns_an_empty_answer():
    p = Scripted(answer(""), answer("Now with text."))
    r = await chat_mod.run_chat(p, {}, [], "hello", [], [])
    assert r.text == "Now with text." and len(p.calls) == 2


async def test_engine_stops_after_max_steps_and_injects_context():
    p = Scripted(*[tool("t")] * 10)
    r = await chat_mod.run_chat(p, {"t": chat_mod.Tool("t", _noop)}, [{"role": "user", "content": "earlier"}], "q",
                                ["Max risk 2% per trade"], [{"filename": "plan.pdf", "kind": "pdf", "summary": "my plan"}],
                                [{"filename": "plan.pdf", "text": "ORB rules"}])
    assert len(p.calls) == chat_mod.MAX_STEPS and r.text  # never loops forever
    system, msgs = p.calls[0]
    assert "Max risk 2% per trade" in system and "plan.pdf" in system and "read-only" in system
    assert msgs[0]["content"] == "earlier" and "ORB rules" in msgs[-1]["content"] and "<attached_file" in msgs[-1]["content"]


async def test_engine_bad_json_is_reported():
    class Bad(Scripted):
        def complete(self, *a):
            return LLMReply("not json")

    with pytest.raises(TranslatorUnavailable):
        await chat_mod.run_chat(Bad(), {}, [], "hi", [], [])


# ---- API -----------------------------------------------------------------------------------------
@pytest.fixture()
def client(tmp_path, monkeypatch):
    s = Settings(_env_file=None, database_url=f"sqlite+aiosqlite:///{(tmp_path / 'a.db').as_posix()}", sim_seed=5,
                 sim_backfill_minutes=0)
    with TestClient(create_app(s)) as c:
        c.provider = Scripted()
        for mod in ("trader.api.routes.chat", "trader.api.chat_tools"):
            monkeypatch.setattr(f"{mod}.build_provider", lambda settings: c.provider)
        yield c


def upload(c, name, data, ctype="text/plain", **q):
    return c.post("/knowledge/upload", params={"filename": name, **q}, content=data, headers={"content-type": ctype})


def test_uploading_the_same_file_again_reuses_the_first_copy(client):
    client.provider.actions = [{"summary": "rules"}]
    a = upload(client, "rules.md", b"Opening range breakout. Max loss 3000.").json()
    b = upload(client, "Rules.MD", b"Opening range breakout. Max loss 3000.").json()
    assert b["id"] == a["id"] and b["duplicate"] is True and "duplicate" not in a
    assert len(client.get("/knowledge/docs").json()["docs"]) == 1
    c = upload(client, "rules.md", b"A different file with the same name.").json()
    assert c["id"] != a["id"] and len(client.get("/knowledge/docs").json()["docs"]) == 2


def test_upload_search_and_delete(client):
    client.provider.actions = [{"summary": "ORB rules and risk limits"}]
    d = upload(client, "orb.md", b"Opening range breakout. Max loss per day is 3000 rupees.").json()
    assert d["status"] == "ready" and d["summary"] == "ORB rules and risk limits" and d["kind"] == "text"
    assert [x["filename"] for x in client.get("/knowledge/docs").json()["docs"]] == ["orb.md"]
    assert upload(client, "x.exe", b"MZ", "application/octet-stream").status_code == 415
    assert upload(client, "e.txt", b"").status_code == 400
    # save=false only reads
    r = upload(client, "peek.txt", b"just reading", save="false").json()
    assert r["text"] == "just reading" and len(client.get("/knowledge/docs").json()["docs"]) == 1
    assert client.delete(f"/knowledge/docs/{d['id']}").status_code == 200
    assert client.delete(f"/knowledge/docs/{d['id']}").status_code == 404


def test_image_is_read_by_the_vision_model(client):
    client.provider.actions = [{"summary": "a chart"}]
    d = upload(client, "chart.png", b"\x89PNG\r\n", "image/png").json()
    assert d["kind"] == "image" and d["status"] == "ready" and client.provider.images == ["image/png"]

    class NoVision(Scripted):
        def describe_image(self, *a):
            raise TranslatorUnavailable("no vision model")

    client.provider = NoVision({"summary": "x"})
    assert upload(client, "c2.png", b"\x89PNG", "image/png").json()["status"] == "no_vision"


def test_conversation_with_tools_memory_and_cards(client):
    client.provider.actions = [{"summary": "risk rules"}]
    upload(client, "rules.txt", b"Never risk more than two percent per trade. Trail the stop after 30 percent profit.")
    conv = client.post("/chat/conversations").json()["id"]
    client.provider.actions = [
        tool("search_knowledge", '{"query":"risk per trade"}'),
        tool("list_strategies"),
        answer("Your notes say max 2% risk per trade.", ["Risk is capped at 2% per trade"]),
    ]
    r = client.post(f"/chat/conversations/{conv}/messages", json={"text": "What is my risk rule?"}).json()
    assert r["content"].startswith("Your notes say") and r["tools_used"] == ["search_knowledge", "list_strategies"]
    assert r["model"] == "fake/m" and r["remembered"] == ["Risk is capped at 2% per trade"]
    hit = client.provider.calls[2][1][-1]["content"]  # calls[0] was the upload summary
    assert "two percent per trade" in hit  # the retrieved passage reached the model as tool data
    notes = client.get("/chat/memory").json()["notes"]
    assert [n["text"] for n in notes] == ["Risk is capped at 2% per trade"]
    # the next turn's prompt carries the memory and the history
    client.provider.actions = [answer("ok")]
    client.post(f"/chat/conversations/{conv}/messages", json={"text": "thanks"})
    system, msgs = client.provider.calls[-1]
    assert "Risk is capped at 2% per trade" in system and msgs[0]["content"] == "What is my risk rule?"
    c = client.get(f"/chat/conversations/{conv}").json()
    assert [m["role"] for m in c["messages"]] == ["user", "assistant"] * 2 and c["title"] == "What is my risk rule?"
    assert client.get("/chat/conversations").json()["conversations"][0]["id"] == conv
    assert client.delete(f"/chat/memory/{notes[0]['id']}").status_code == 200
    assert client.delete(f"/chat/conversations/{conv}").status_code == 200
    assert client.get(f"/chat/conversations/{conv}").status_code == 404


def test_at_mentioned_file_limits_the_search_to_that_file(client):
    client.provider.actions = [{"summary": "a"}, {"summary": "b"}]
    a = upload(client, "alpha.txt", b"Risk rule from alpha: stop loss is 10 percent.").json()
    upload(client, "beta.txt", b"Risk rule from beta: stop loss is 40 percent.")
    conv = client.post("/chat/conversations").json()["id"]
    client.provider.actions = [tool("search_knowledge", '{"query":"risk rule stop loss"}'), answer("done")]
    client.post(f"/chat/conversations/{conv}/messages", json={"text": "What does @alpha.txt say about risk?", "focus_doc_ids": [a["id"]]})
    system, msgs = client.provider.calls[-1]
    first_user = client.provider.calls[-2][1][-1]["content"]
    assert "referred to these files with @" in first_user and "alpha.txt" in first_user
    hit = msgs[-1]["content"]
    assert "alpha" in hit and "beta" not in hit  # beta.txt was never searched


def test_attached_file_text_reaches_the_model_and_memory_dedupes(client):
    client.provider.actions = [{"summary": "s"}]
    doc = upload(client, "plan.txt", b"MY SECRET PLAN: buy dips on NIFTY").json()
    conv = client.post("/chat/conversations").json()["id"]
    client.provider.actions = [answer("Read it.", ["Likes dip buying", "likes dip buying"])]
    client.post(f"/chat/conversations/{conv}/messages", json={"text": "summarise", "doc_ids": [doc["id"]]})
    assert "MY SECRET PLAN" in client.provider.calls[0][1][-1]["content"]
    assert len(client.get("/chat/memory").json()["notes"]) == 1


def test_draft_strategy_tool_returns_a_card_and_saves_nothing(client):
    from trader.strategy.spec import MOMENTUM_V1

    env = {"spec_json": json.dumps(MOMENTUM_V1), "assumptions": ["RSI 14"], "questions": [], "unsupported": []}
    conv = client.post("/chat/conversations").json()["id"]
    # the same fake serves the chat turns and the translator call inside the tool, in call order
    client.provider.actions = [tool("draft_strategy", '{"description":"momentum on nifty 5 min with RSI"}'), env,
                               answer("Draft ready, review and save it.")]
    r = client.post(f"/chat/conversations/{conv}/messages", json={"text": "build a momentum strategy"}).json()
    assert r["cards"][0]["type"] == "strategy_draft" and r["cards"][0]["spec"]["name"] == MOMENTUM_V1["name"]
    assert client.get("/strategies").json() == []  # nothing saved until the user presses Save


def test_provider_outage_is_a_clear_503_and_stores_nothing(client):
    conv = client.post("/chat/conversations").json()["id"]

    class Down(Scripted):
        def complete(self, *a):
            raise TranslatorUnavailable("groq is rate limiting this key")

    client.provider = Down()
    r = client.post(f"/chat/conversations/{conv}/messages", json={"text": "hello there"})
    assert r.status_code == 503 and "rate limiting" in r.json()["detail"]
    assert client.get(f"/chat/conversations/{conv}").json()["messages"] == []


def test_assistant_has_no_way_to_trade():
    """Read-only by construction: the tool set is fixed, and none of it touches orders or the runner."""
    import inspect
    from types import SimpleNamespace

    from trader.api import chat_tools

    names = {"list_strategies", "get_strategy", "list_backtests", "get_backtest", "paper_journal", "paper_status",
             "search_knowledge", "draft_strategy"}
    from trader.accounts.repo import local_principal

    app = SimpleNamespace(state=SimpleNamespace(sessionmaker=None, settings=None, md=SimpleNamespace(settings=None)))
    assert set(chat_tools.build_tools(SimpleNamespace(app=app), local_principal(), None)) == names
    src = inspect.getsource(chat_tools)
    for banned in ("submit", "place_order", "broker.", ".start(", ".kill(", "trip_kill_switch", "RiskEngine"):
        assert banned not in src, banned


async def test_engine_looks_up_data_questions_itself_so_one_model_call_is_enough():
    async def ls(_):
        return [{"name": "X"}]

    p = Scripted(answer("You have 1 strategy: X."))
    r = await chat_mod.run_chat(p, {"list_strategies": chat_mod.Tool("l", ls)}, [], "list my saved strategies", [], [],
                                snapshot="Saved strategies (1):\n  - X (v1, 5m)")
    assert r.text == "You have 1 strategy: X." and r.tools_used == ["list_strategies"] and len(p.calls) == 1
    assert '"name":"X"' in p.calls[0][1][-1]["content"] and "<already_looked_up>" in p.calls[0][1][-1]["content"]
    assert "Saved strategies (1)" in p.calls[0][0]  # real facts are in the prompt


async def test_engine_does_not_nudge_small_talk():
    p = Scripted(answer("Hello!"))
    r = await chat_mod.run_chat(p, {"t": chat_mod.Tool("t", _noop)}, [], "hello there", [], [])
    assert r.text == "Hello!" and len(p.calls) == 1


async def test_engine_prefetches_file_search_with_the_question_and_runs_lookups_together():
    fetched = []

    async def search(a):
        fetched.append(a)
        return [{"file": "rules.txt", "text": "Never risk more than 2 percent."}]

    p = Scripted(answer("Your notes say 2% per trade."))
    r = await chat_mod.run_chat(p, {"search_knowledge": chat_mod.Tool("s", search)}, [], "what do my notes say about risk", [], [])
    assert fetched == [{"query": "what do my notes say about risk"}] and r.tools_used == ["search_knowledge"]
    assert r.text == "Your notes say 2% per trade." and len(p.calls) == 1
    assert "Never risk more than 2 percent" in p.calls[0][1][-1]["content"]


def test_history_keeps_recent_turns_in_full_and_shortens_older_replies():
    long = "x" * 3000
    hist = [{"role": "user", "content": "q"}, {"role": "assistant", "content": long}] * 6
    out = chat_mod._trim_history(hist)
    assert len(out) == chat_mod.MAX_HISTORY
    assert len(out[0]["content"]) <= chat_mod.OLD_REPLY_CHARS + 1 or out[0]["role"] == "user"
    assert out[-1]["content"] == long  # the latest reply is untouched


# ---- talking about the strategy open on the Strategies page ----------------------------------
def test_strategy_context_is_built_from_the_rules_and_the_owners_own_backtest():
    from datetime import date, datetime, timezone

    from trader.api import strategy_context as sc
    from trader.db.models import BacktestRunRow
    from trader.strategy.spec import MOMENTUM_V1

    ctx = sc.StrategyContextIn(spec=MOMENTUM_V1, backtest_id=None)
    text = sc.build(ctx, None)
    assert "NOT been backtested" in text and "draft_strategy" in text and "Buy the ATM call" in text
    run = BacktestRunRow(id="r1", owner="u", created_at=datetime.now(timezone.utc), underlying="NIFTY", start=date(2026, 1, 1),
                         end=date(2026, 6, 1), timeframe="5m", strategy="x", data_source="real", pricing_mode="synthetic",
                         result={"metrics": {"trades": 12, "total_net_pnl": -500, "win_rate": 33.3, "exit_reasons": {"stop loss": 8}},
                                 "trades": [{"contract": "C", "entry_time": "2026-02-02T04:00:00", "exit_reason": "stop loss", "net_pnl": -900}],
                                 "params": {"strategy": {}}, "warnings": ["w"]})
    text = sc.build(ctx, run)
    assert "id r1" in text and "net P&L -500" in text and "stop loss" in text and "worst trades" in text
    assert "not valid yet" in sc.build(sc.StrategyContextIn(spec={"nope": 1}), None)


def test_talking_about_a_strategy_sends_the_rules_but_stores_only_the_question(client):
    from trader.strategy.spec import MOMENTUM_V1

    conv = client.post("/chat/conversations").json()["id"]
    client.provider.actions = [answer("Tighten the stop and add an ADX filter.")]
    r = client.post(f"/chat/conversations/{conv}/messages",
                    json={"text": "How can I improve this?", "context": {"spec": MOMENTUM_V1, "backtest_id": "someone-elses"}})
    assert r.status_code == 200 and r.json()["content"].startswith("Tighten")
    sent = client.provider.calls[0][1][-1]["content"]
    assert "Buy the ATM call" in sent and "How can I improve this?" in sent and "NOT been backtested" in sent
    assert len(client.provider.calls) == 1  # no look-up round trip: the context already has the data
    stored = client.get(f"/chat/conversations/{conv}").json()["messages"]
    assert stored[0]["content"] == "How can I improve this?"  # the long context is not kept in the history
